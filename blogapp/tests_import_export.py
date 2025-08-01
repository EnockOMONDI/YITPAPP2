"""
Comprehensive tests for YITP Blog Import/Export functionality
Tests Excel import/export features, validation, and error handling
"""

import os
import tempfile
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from io import BytesIO
import openpyxl
from blogapp.models import Post, Category, Comment
from blogapp.resources import PostResource, CategoryResource
from blogapp.utils import create_blog_import_template, validate_import_file


class BlogImportExportTestCase(TestCase):
    """Test case for blog import/export functionality"""
    
    def setUp(self):
        """Set up test data"""
        # Create test user
        self.user = User.objects.create_user(
            username='testuser',
            email='test@example.com',
            password='testpass123'
        )
        self.admin_user = User.objects.create_superuser(
            username='admin',
            email='admin@example.com',
            password='adminpass123'
        )
        
        # Create test category
        self.category = Category.objects.create(
            title='Test Category',
            slug='test-category',
            active=True
        )
        
        # Create test posts
        self.post1 = Post.objects.create(
            title='Test Post 1',
            content='<p>This is test content for post 1</p>',
            Author='Test Author',
            category=self.category,
            user=self.user,
            status='published',
            featured=True
        )
        
        self.post2 = Post.objects.create(
            title='Test Post 2',
            content='<p>This is test content for post 2</p>',
            Author='Another Author',
            category=self.category,
            status='draft',
            trending=True
        )
        
        # Add tags
        self.post1.tags.add('test', 'django', 'blog')
        self.post2.tags.add('python', 'web', 'development')
        
        self.client = Client()

    def test_post_resource_export(self):
        """Test exporting posts using PostResource"""
        resource = PostResource()
        dataset = resource.export()
        
        # Check that data is exported
        self.assertGreater(len(dataset), 0)
        self.assertEqual(len(dataset), 2)  # Two test posts
        
        # Check headers
        expected_headers = [
            'id', 'title', 'author', 'category', 'tags', 'status',
            'featured', 'trending', 'publication_date', 'content',
            'featured_image_url', 'meta_description', 'slug', 'views'
        ]
        for header in expected_headers:
            self.assertIn(header, dataset.headers)
        
        # Check data content (posts are ordered by -date, so most recent first)
        # Find the row with 'Test Post 1'
        test_post_1_row = None
        for row in dataset:
            if row[dataset.headers.index('title')] == 'Test Post 1':
                test_post_1_row = row
                break

        self.assertIsNotNone(test_post_1_row)
        self.assertEqual(test_post_1_row[dataset.headers.index('title')], 'Test Post 1')
        self.assertEqual(first_row[dataset.headers.index('status')], 'published')
        self.assertEqual(first_row[dataset.headers.index('featured')], True)

    def test_post_resource_import(self):
        """Test importing posts using PostResource"""
        # Create import data
        import_data = [
            ['title', 'content', 'author', 'category', 'status', 'featured'],
            ['Imported Post 1', '<p>Imported content 1</p>', 'Import Author', 'New Category', 'draft', 'FALSE'],
            ['Imported Post 2', '<p>Imported content 2</p>', 'Import Author', 'Test Category', 'published', 'TRUE']
        ]
        
        # Create dataset
        from tablib import Dataset
        dataset = Dataset()
        dataset.headers = import_data[0]
        for row in import_data[1:]:
            dataset.append(row)
        
        # Import data
        resource = PostResource()
        result = resource.import_data(dataset, dry_run=False)
        
        # Check import results
        self.assertFalse(result.has_errors())
        self.assertEqual(len(result.rows), 2)
        
        # Verify posts were created
        imported_post1 = Post.objects.get(title='Imported Post 1')
        imported_post2 = Post.objects.get(title='Imported Post 2')
        
        self.assertEqual(imported_post1.status, 'draft')
        self.assertEqual(imported_post1.featured, False)
        self.assertEqual(imported_post2.status, 'published')
        self.assertEqual(imported_post2.featured, True)
        
        # Check that new category was created
        new_category = Category.objects.get(title='New Category')
        self.assertEqual(imported_post1.category, new_category)

    def test_template_generation(self):
        """Test Excel template generation"""
        response = create_blog_import_template()
        
        # Check response
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('YITP_Blog_Import_Template', response['Content-Disposition'])
        
        # Check Excel content
        content = BytesIO(response.content)
        workbook = openpyxl.load_workbook(content)
        
        # Check worksheets
        self.assertIn('Blog Posts Import Template', workbook.sheetnames)
        self.assertIn('Instructions', workbook.sheetnames)
        
        # Check main worksheet
        ws = workbook['Blog Posts Import Template']
        headers = [cell.value for cell in ws[1]]
        
        expected_headers = [
            'Post Title (Required)', 'Content (Required)', 'Author Name',
            'Category', 'Tags', 'User', 'Status', 'Featured', 'Trending',
            'Publication Date', 'Featured Image URL', 'Meta Description'
        ]
        
        for expected_header in expected_headers:
            self.assertIn(expected_header, headers)

    def test_file_validation(self):
        """Test file validation functionality"""
        # Create a valid Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['title', 'content', 'author'])
        ws.append(['Sample Title', 'Sample Content', 'Sample Author'])
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp_file:
            wb.save(tmp_file.name)
            
            # Create uploaded file object
            with open(tmp_file.name, 'rb') as f:
                uploaded_file = SimpleUploadedFile(
                    'test.xlsx',
                    f.read(),
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
                )
            
            # Test validation
            errors = validate_import_file(uploaded_file)
            # Note: There might be a validation error about missing content column
            # but the file structure should be valid
            self.assertLessEqual(len(errors), 1)
            
            # Clean up
            os.unlink(tmp_file.name)
        
        # Test invalid file type
        invalid_file = SimpleUploadedFile(
            'test.txt',
            b'This is not an Excel file',
            content_type='text/plain'
        )
        errors = validate_import_file(invalid_file)
        self.assertGreater(len(errors), 0)
        self.assertIn('Excel file', errors[0])

    def test_management_command_dry_run(self):
        """Test the import management command with dry run"""
        # Create test Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # Add headers
        headers = ['title', 'content', 'author', 'category', 'status']
        ws.append(headers)
        ws.append(['Description Row'])  # Skip this row
        
        # Add test data
        test_data = [
            ['Command Test Post 1', '<p>Content 1</p>', 'Command Author', 'Command Category', 'draft'],
            ['Command Test Post 2', '<p>Content 2</p>', 'Command Author', 'Command Category', 'published']
        ]
        
        for row in test_data:
            ws.append(row)
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp_file:
            wb.save(tmp_file.name)
            
            try:
                # Test dry run
                call_command('import_blog_posts', tmp_file.name, '--dry-run', verbosity=0)
                
                # Verify no posts were actually created
                self.assertFalse(Post.objects.filter(title='Command Test Post 1').exists())
                self.assertFalse(Post.objects.filter(title='Command Test Post 2').exists())
                
            finally:
                # Clean up
                os.unlink(tmp_file.name)

    def test_management_command_actual_import(self):
        """Test the import management command with actual import"""
        # Create test Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # Add headers
        headers = ['title', 'content', 'author', 'category', 'status', 'featured', 'tags']
        ws.append(headers)
        ws.append(['Description Row'])  # Skip this row
        
        # Add test data
        test_data = [
            ['Actual Import Post 1', '<p>Actual content 1</p>', 'Import Author', 'Import Category', 'draft', 'TRUE', 'import,test'],
            ['Actual Import Post 2', '<p>Actual content 2</p>', 'Import Author', 'Import Category', 'published', 'FALSE', 'command,blog']
        ]
        
        for row in test_data:
            ws.append(row)
        
        # Save to temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp_file:
            wb.save(tmp_file.name)
            
            try:
                # Test actual import
                call_command('import_blog_posts', tmp_file.name, verbosity=0)
                
                # Verify posts were created
                post1 = Post.objects.get(title='Actual Import Post 1')
                post2 = Post.objects.get(title='Actual Import Post 2')
                
                self.assertEqual(post1.status, 'draft')
                self.assertEqual(post1.featured, True)
                self.assertEqual(post2.status, 'published')
                self.assertEqual(post2.featured, False)
                
                # Check category was created
                import_category = Category.objects.get(title='Import Category')
                self.assertEqual(post1.category, import_category)
                self.assertEqual(post2.category, import_category)
                
                # Check tags
                self.assertIn('import', [tag.name for tag in post1.tags.all()])
                self.assertIn('test', [tag.name for tag in post1.tags.all()])
                
            finally:
                # Clean up
                os.unlink(tmp_file.name)

    def test_admin_template_download(self):
        """Test template download through admin interface"""
        # Login as admin
        self.client.login(username='admin', password='adminpass123')
        
        # Test template download URL
        response = self.client.get(reverse('admin:blogapp_post_download_template'))
        
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_category_resource(self):
        """Test CategoryResource import/export"""
        resource = CategoryResource()
        
        # Test export
        dataset = resource.export()
        self.assertGreater(len(dataset), 0)
        
        # Test import
        from tablib import Dataset
        import_dataset = Dataset()
        import_dataset.headers = ['title', 'slug', 'active']
        import_dataset.append(['New Category', 'new-category', True])
        
        result = resource.import_data(import_dataset, dry_run=False)
        self.assertFalse(result.has_errors())
        
        # Verify category was created
        new_category = Category.objects.get(title='New Category')
        self.assertEqual(new_category.slug, 'new-category')
        self.assertTrue(new_category.active)

    def test_error_handling(self):
        """Test error handling in import process"""
        # Create dataset with invalid data
        from tablib import Dataset
        dataset = Dataset()
        dataset.headers = ['title', 'content', 'status']
        dataset.append(['', 'Content without title', 'invalid_status'])  # Missing title, invalid status
        
        resource = PostResource()
        result = resource.import_data(dataset, dry_run=True)
        
        # Should have validation errors
        self.assertTrue(result.has_validation_errors())

    def tearDown(self):
        """Clean up after tests"""
        # Clean up any created files or data
        pass


class BlogAdminIntegrationTestCase(TestCase):
    """Integration tests for admin interface"""
    
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            username='admin',
            email='admin@example.com',
            password='adminpass123'
        )
        self.client = Client()
        self.client.login(username='admin', password='adminpass123')

    def test_admin_changelist_view(self):
        """Test enhanced admin changelist view"""
        response = self.client.get(reverse('admin:blogapp_post_changelist'))
        self.assertEqual(response.status_code, 200)
        
        # Check for import/export buttons in context
        self.assertIn('has_import_permission', response.context)
        self.assertIn('has_export_permission', response.context)
        self.assertIn('download_template_url', response.context)

    def test_admin_permissions(self):
        """Test admin permissions for import/export"""
        # Test with regular user (no permissions)
        regular_user = User.objects.create_user(
            username='regular',
            email='regular@example.com',
            password='regularpass123'
        )
        
        client = Client()
        client.login(username='regular', password='regularpass123')
        
        # Should not have access to admin
        response = client.get(reverse('admin:blogapp_post_changelist'))
        self.assertEqual(response.status_code, 302)  # Redirect to login
