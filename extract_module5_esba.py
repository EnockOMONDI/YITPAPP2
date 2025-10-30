#!/usr/bin/env python3
"""
Module 5: Entrepreneurship & Small Business Administration (ESBA) Content Extraction
Extracts 4 lessons from the Word document following the mind map structure
"""

import os
import sys
import django
import json
from datetime import datetime
import docx
from docx import Document
import html
import re

# Setup Django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'blog.settings')
django.setup()

def extract_lesson_content(doc, start_para, end_para, lesson_title):
    """Extract content from Word document for a specific lesson"""
    print(f"\n=== Extracting {lesson_title} ===")
    print(f"Paragraphs: {start_para} - {end_para}")
    
    content_html = []
    
    # Add lesson introduction with YITP branding
    content_html.append(f'''
    <div class="lesson-header" style="background: linear-gradient(135deg, #1a2e53 0%, #ff5d15 100%); color: white; padding: 30px; border-radius: 10px; margin-bottom: 30px;">
        <h1 style="color: white; margin: 0; font-size: 2.2rem; font-weight: bold;">{lesson_title}</h1>
        <p style="color: #f8f9fa; margin: 10px 0 0 0; font-size: 1.1rem;">Youth Impact Training Programme - Module 5: ESBA</p>
    </div>
    ''')
    
    current_section = ""
    section_content = []
    
    for i in range(start_para - 1, min(end_para, len(doc.paragraphs))):
        para = doc.paragraphs[i]
        text = para.text.strip()
        
        if not text:
            continue
            
        # Check if this is a major heading
        is_major_heading = (
            text.isupper() and len(text) < 100 and 
            any(keyword in text.upper() for keyword in [
                'STARTING A BUSINESS', 'BUILDING', 'ACTUALISING', 
                'BUSINESS MANAGEMENT', 'SUSTAINABILITY', 'IDEATION', 
                'BUSINESS PLAN', 'PREMISES', 'CAPITAL'
            ])
        )
        
        # Check if this is a section heading
        is_section_heading = (
            (text.endswith(':') and len(text) < 150) or
            (text.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.')) and len(text) < 150) or
            any(keyword in text for keyword in ['Importance:', 'Definition:', 'Examples:'])
        )
        
        if is_major_heading:
            # Save previous section if exists
            if section_content:
                content_html.extend(section_content)
                section_content = []
            
            # Start new major section
            current_section = text
            content_html.append(f'''
            <div class="major-section" style="margin: 40px 0 30px 0;">
                <h2 style="color: #1a2e53; font-size: 1.8rem; font-weight: bold; border-bottom: 3px solid #ff5d15; padding-bottom: 10px; margin-bottom: 20px;">
                    {html.escape(text)}
                </h2>
            </div>
            ''')
            
        elif is_section_heading:
            # Add section heading
            section_content.append(f'''
            <div class="section-heading" style="margin: 25px 0 15px 0;">
                <h3 style="color: #ff5d15; font-size: 1.4rem; font-weight: bold; margin-bottom: 10px;">
                    {html.escape(text)}
                </h3>
            </div>
            ''')
            
        else:
            # Regular content paragraph
            # Check if it's a list item
            if text.startswith(('•', '-', '→', '✓')) or re.match(r'^\d+\.', text):
                section_content.append(f'''
                <div class="list-item" style="margin: 8px 0; padding-left: 20px; border-left: 3px solid #ff5d15;">
                    <p style="margin: 0; color: #333; line-height: 1.6;">
                        {html.escape(text)}
                    </p>
                </div>
                ''')
            else:
                # Regular paragraph
                section_content.append(f'''
                <div class="content-paragraph" style="margin: 15px 0;">
                    <p style="color: #333; line-height: 1.7; font-size: 1rem; text-align: justify;">
                        {html.escape(text)}
                    </p>
                </div>
                ''')
    
    # Add remaining section content
    if section_content:
        content_html.extend(section_content)
    
    # Add lesson footer
    content_html.append(f'''
    <div class="lesson-footer" style="margin-top: 40px; padding: 20px; background-color: #f8f9fa; border-radius: 8px; border-left: 5px solid #ff5d15;">
        <p style="margin: 0; color: #1a2e53; font-weight: bold;">
            🎯 Lesson Complete: {lesson_title}
        </p>
        <p style="margin: 5px 0 0 0; color: #666; font-size: 0.9rem;">
            Continue to the next lesson to build on these concepts.
        </p>
    </div>
    ''')
    
    return ''.join(content_html)

def generate_quiz_questions(lesson_number, lesson_title, content_preview):
    """Generate 5 quiz questions for each lesson"""
    
    quiz_templates = {
        1: [  # Starting a Business
            {
                "question": "What is the primary purpose of conducting a SWOT analysis for a business idea?",
                "options": [
                    "To determine the business location",
                    "To evaluate strengths, weaknesses, opportunities, and threats",
                    "To calculate startup costs",
                    "To choose a business name"
                ],
                "correct_answer": 1,
                "explanation": "SWOT analysis helps entrepreneurs evaluate the internal strengths and weaknesses of their business idea, as well as external opportunities and threats in the market."
            },
            {
                "question": "Which section of a business plan provides a detailed explanation of your sales strategy and pricing plan?",
                "options": [
                    "Executive Summary",
                    "Market Analysis",
                    "Sales and Marketing Plan",
                    "Financial Plan"
                ],
                "correct_answer": 2,
                "explanation": "The Sales and Marketing Plan section outlines your sales strategy, pricing plan, and how you will attract and retain customers."
            },
            {
                "question": "What is a key advantage of a Limited Liability Company (LLC) business structure?",
                "options": [
                    "Unlimited personal liability",
                    "Complex tax requirements",
                    "Personal asset protection from business debts",
                    "Requires multiple owners"
                ],
                "correct_answer": 2,
                "explanation": "LLCs provide personal asset protection, meaning owners' personal assets are generally protected from business debts and liabilities."
            },
            {
                "question": "Which financial statement provides a snapshot of a company's assets, liabilities, and equity at a specific point in time?",
                "options": [
                    "Income Statement",
                    "Cash Flow Statement",
                    "Balance Sheet",
                    "Profit and Loss Statement"
                ],
                "correct_answer": 2,
                "explanation": "The Balance Sheet shows the company's financial position at a specific point in time, listing all assets, liabilities, and owner's equity."
            },
            {
                "question": "What should be the first step when developing a business idea?",
                "options": [
                    "Register the business name",
                    "Identify a market need or gap",
                    "Secure funding",
                    "Hire employees"
                ],
                "correct_answer": 1,
                "explanation": "Identifying a market need or gap is crucial as it ensures your business idea addresses a real problem or demand in the market."
            }
        ],
        2: [  # Building & Actualizing Business
            {
                "question": "When choosing business premises, which factor is most important for a retail business?",
                "options": [
                    "Low rent costs only",
                    "Proximity to target customers and foot traffic",
                    "Large storage space",
                    "Availability of parking for employees"
                ],
                "correct_answer": 1,
                "explanation": "For retail businesses, location with high foot traffic and proximity to target customers is crucial for success and revenue generation."
            },
            {
                "question": "What is a key characteristic of venture capital funding?",
                "options": [
                    "No involvement in business decisions",
                    "Focus on low-growth, stable businesses",
                    "Investment in high-growth potential companies with active involvement",
                    "Only provides small amounts of funding"
                ],
                "correct_answer": 2,
                "explanation": "Venture capitalists invest in high-growth potential companies and typically want to play an active role in the companies they finance."
            },
            {
                "question": "Which of the following is NOT typically required for business registration in Kenya?",
                "options": [
                    "Certificate of Registration (CR1)",
                    "Memorandum and Articles of Association",
                    "Personal bank statements of all employees",
                    "KRA PIN registration"
                ],
                "correct_answer": 2,
                "explanation": "Personal bank statements of employees are not required for business registration. The focus is on business documentation and tax registration."
            },
            {
                "question": "What is the primary advantage of hiring friends and family members in a startup?",
                "options": [
                    "They always have the required skills",
                    "Cost-effective way to build a loyal team",
                    "No need for formal contracts",
                    "They work for free permanently"
                ],
                "correct_answer": 1,
                "explanation": "Hiring friends and family can be cost-effective and help build a loyal team that believes in your vision, though proper agreements are still important."
            },
            {
                "question": "What does OSHA registration ensure for a business workplace?",
                "options": [
                    "Higher employee salaries",
                    "Workplace is free of hazards and complies with safety standards",
                    "Automatic business insurance coverage",
                    "Exemption from other licenses"
                ],
                "correct_answer": 1,
                "explanation": "OSHA (Occupational Safety and Health Services) registration ensures the workplace meets established safety standards and is free of hazards for employees."
            }
        ],
        3: [  # Business Management
            {
                "question": "What is the primary purpose of strategic planning in business management?",
                "options": [
                    "To manage daily operations",
                    "To set long-term goals and define steps to achieve them",
                    "To handle customer complaints",
                    "To calculate monthly expenses"
                ],
                "correct_answer": 1,
                "explanation": "Strategic planning involves setting long-term goals and defining the steps needed to achieve them, providing direction for the business."
            },
            {
                "question": "Which of the following is a common cash flow problem faced by startups?",
                "options": [
                    "Too much profit",
                    "Late payments by customers",
                    "Excessive employee satisfaction",
                    "Over-investment in marketing"
                ],
                "correct_answer": 1,
                "explanation": "Late-paying customers can severely impact a startup's cash flow, creating dangerous financial situations."
            },
            {
                "question": "What is the main benefit of effective marketing and sales strategies?",
                "options": [
                    "Reducing operational costs",
                    "Attracting and retaining customers",
                    "Eliminating competition",
                    "Avoiding regulatory compliance"
                ],
                "correct_answer": 1,
                "explanation": "Effective marketing and sales strategies are essential for attracting new customers and retaining existing ones, driving business growth."
            },
            {
                "question": "In business intelligence (BI), what is the primary advantage of data-driven decision making?",
                "options": [
                    "It eliminates the need for human judgment",
                    "It provides faster, more accurate decision-making",
                    "It reduces the need for technology",
                    "It guarantees business success"
                ],
                "correct_answer": 1,
                "explanation": "Business intelligence tools enable faster, more accurate decision-making by providing data-driven insights into business operations and customer behavior."
            },
            {
                "question": "What is a key component of basic risk and crisis management?",
                "options": [
                    "Ignoring potential threats",
                    "Proactively identifying and mitigating potential risks",
                    "Waiting for crises to occur before responding",
                    "Focusing only on financial risks"
                ],
                "correct_answer": 1,
                "explanation": "Basic risk and crisis management involves proactively identifying potential threats and preparing contingency plans to mitigate them."
            }
        ],
        4: [  # Business Sustainability
            {
                "question": "What defines a sustainable business according to the module?",
                "options": [
                    "A business that only focuses on profit maximization",
                    "A business that can support itself, offer profit, and balance consumer/employee needs",
                    "A business that operates without any environmental impact",
                    "A business that never changes its operations"
                ],
                "correct_answer": 1,
                "explanation": "A sustainable business can support itself financially, provide profit to owners, while balancing the needs of both consumers and employees."
            },
            {
                "question": "Which factor is crucial for building a sustainable business foundation?",
                "options": [
                    "Cutting costs at all expenses",
                    "Building on tried-and-tested business principles",
                    "Avoiding all forms of competition",
                    "Focusing only on short-term gains"
                ],
                "correct_answer": 1,
                "explanation": "A strong foundation built on tried-and-tested business principles is crucial for developing real sustainability in business."
            },
            {
                "question": "Why is customer experience important for business sustainability?",
                "options": [
                    "It reduces operational costs",
                    "It eliminates the need for marketing",
                    "It helps build customer loyalty and sustain business during economic challenges",
                    "It guarantees immediate profits"
                ],
                "correct_answer": 2,
                "explanation": "Superior customer experience builds loyalty, which is critical for sustaining business even during economic downturns or challenging periods."
            },
            {
                "question": "What role does technology play in sustainable business growth?",
                "options": [
                    "It replaces the need for human employees",
                    "It enables data-driven decision making for sustainable growth",
                    "It eliminates all business risks",
                    "It guarantees immediate success"
                ],
                "correct_answer": 1,
                "explanation": "Technology and data-driven decision making help ensure sustainable growth by providing insights for informed business decisions."
            },
            {
                "question": "According to the module, what is essential for maintaining business integrity?",
                "options": [
                    "Maximizing profits regardless of methods",
                    "Avoiding all business partnerships",
                    "Never getting complacent and maintaining operational excellence",
                    "Focusing only on internal operations"
                ],
                "correct_answer": 2,
                "explanation": "Never getting complacent and maintaining perfectionism in operational execution is essential for building sustainable business integrity."
            }
        ]
    }

    return quiz_templates.get(lesson_number, [])

def main():
    print("=== MODULE 5: ESBA CONTENT EXTRACTION ===")
    print("Extracting 4 lessons from Word document...")
    
    try:
        # Load the Word document
        doc_path = 'courseunits/5. ESBA .docx'
        doc = Document(doc_path)
        print(f"Loaded document: {doc_path}")
        print(f"Total paragraphs: {len(doc.paragraphs)}")
        
        # Define lesson boundaries based on content analysis
        lessons = [
            {
                "title": "Starting a Business",
                "start_para": 27,
                "end_para": 144,
                "sections": ["Ideation", "Business Plan", "Business Structures", "Financial Statements"],
                "duration": 85
            },
            {
                "title": "Building & Actualizing Business", 
                "start_para": 145,
                "end_para": 307,
                "sections": ["Premises", "Capital (Sources of Financing)", "Skills & Labour", "Licensing"],
                "duration": 95
            },
            {
                "title": "Business Management",
                "start_para": 308,
                "end_para": 464,
                "sections": ["Strategic Planning", "Financial Management", "Marketing & Sales", "Operations", "HR", "Technology", "Communication", "Business Intelligence", "Risk Management"],
                "duration": 90
            },
            {
                "title": "Business Sustainability",
                "start_para": 465,
                "end_para": 508,
                "sections": ["Definition & Importance", "Achieving Sustainability"],
                "duration": 55
            }
        ]
        
        # Extract content for each lesson
        extracted_lessons = []
        
        for i, lesson_info in enumerate(lessons, 1):
            lesson_title = f"Lesson {i}: {lesson_info['title']}"
            
            # Extract lesson content
            content_html = extract_lesson_content(
                doc, 
                lesson_info['start_para'], 
                lesson_info['end_para'], 
                lesson_title
            )
            
            # Generate quiz questions
            quiz_questions = generate_quiz_questions(i, lesson_title, content_html[:500])
            
            lesson_data = {
                "lesson_number": i,
                "title": lesson_title,
                "content": content_html,
                "sections": lesson_info['sections'],
                "estimated_duration": lesson_info['duration'],
                "quiz_questions": quiz_questions,
                "extracted_at": datetime.now().isoformat()
            }
            
            extracted_lessons.append(lesson_data)
            print(f"✅ Extracted: {lesson_title} ({len(content_html)} characters)")
        
        print(f"\n=== EXTRACTION COMPLETE ===")
        print(f"Total lessons extracted: {len(extracted_lessons)}")
        
        # Save extracted content to JSON file for review
        output_file = 'module5_esba_extracted.json'
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(extracted_lessons, f, indent=2, ensure_ascii=False)
        
        print(f"✅ Content saved to: {output_file}")
        print("\nNext step: Review the extracted content and run import script")
        
        return extracted_lessons
        
    except Exception as e:
        print(f"❌ Error during extraction: {e}")
        import traceback
        traceback.print_exc()
        return None

if __name__ == "__main__":
    main()
